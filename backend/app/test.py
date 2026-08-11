from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def add_bottom_border(paragraph):
    pPr = paragraph._p.get_or_add_pPr()

    pBdr = OxmlElement("w:pBdr")

    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), "000000")

    pBdr.append(bottom)
    pPr.append(pBdr)


doc = Document()

p = doc.add_paragraph()

run = p.add_run("SKILLS")
run.bold = True
run.font.name = "Calibri"
run.font.size = Pt(13)
run.font.color.rgb = RGBColor(0, 0, 0)

add_bottom_border(p)

doc.add_paragraph("Python, FastAPI, PostgreSQL, Docker")

doc.save("border_test.docx")

print("border_test.docx created")